"""Tests for Excel reading/writing safety and the end-to-end write plan.

A workbook is built in-memory with arbitrary sheet/column/date values (none of
which appear in application code) to verify: column selection by letter,
repeated-header handling, formula preservation, no-overwrite-without-confirm,
and audit generation.
"""
import datetime as dt
import io

import openpyxl
from openpyxl.utils import column_index_from_string

from src import excel_reader, excel_writer
from src.aggregator import aggregate
from src.audit import build_audit_records
from src.excel_writer import WriteAction, WriteOp, apply_writes
from src.mapping import (
    AppConfig,
    OutputType,
    SheetMapping,
    SourceMapping,
    TargetMode,
    WriteRules,
)
from src.parser import parse_text_lines
from src.validator import Status, build_plan, plan_to_write_ops


def _make_workbook_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S1"
    # Row 1 headers include a REPEATED header name to prove we don't rely on it.
    ws.append(["Date", "Amount", "Amount", "Note"])  # two "Amount" columns
    ws.append([dt.date(2026, 5, 15), None, None, "row1"])
    ws.append([dt.date(2026, 5, 16), None, None, "row2"])
    # A formula in an untouched cell to verify preservation.
    ws["E2"] = "=B2+C2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_resolve_column_by_letter_and_repeated_header():
    data = _make_workbook_bytes()
    wb = excel_reader.open_workbook(data)
    # Header name "Amount" is duplicated; detection must surface it.
    dups = excel_reader.find_duplicate_headers(wb, "S1", 1)
    assert "Amount" in dups
    # We can still target the SECOND amount column unambiguously by letter "C".
    idx = excel_reader.resolve_column_index(wb, "S1", mode="column_letter", selector="C", header_row=1)
    assert idx == column_index_from_string("C")


def test_find_row_by_date_real_dateobject():
    data = _make_workbook_bytes()
    wb = excel_reader.open_workbook(data)
    row = excel_reader.find_row_by_date(wb, "S1", 1, dt.date(2026, 5, 16), header_row=1)
    assert row == 3


def _config_for_letter_target():
    return AppConfig(
        source=SourceMapping(
            date_field="Inv Date", amount_field="Amount", sum_duplicate_dates=True,
            decimal_sep=".", thousands_sep=",",
        ),
        group_to_sheet={"SOME GROUP": "S1"},
        sheets={"S1": SheetMapping(
            sheet_name="S1", header_row=1,
            date_target_mode=TargetMode.COLUMN_LETTER, date_column="A",
            amount_target_mode=TargetMode.COLUMN_LETTER, amount_column="C",
        )},
        write_rules=WriteRules(write_action=WriteAction.REPLACE, output_type=OutputType.NUMERIC,
                               aggregation_keys=["group", "date"]),
    )


def _records():
    text = """\
Customer Name: SOME GROUP
A1   15/05/2026   630.00
A2   15/05/2026   12,705.00
"""
    return parse_text_lines([text], group_label="Customer Name",
                            field_names=["Inv No", "Inv Date", "Amount"])


def test_full_plan_and_write_preserves_formula():
    data = _make_workbook_bytes()
    cfg = _config_for_letter_target()
    records = _records()
    rows = aggregate(records, cfg.source, aggregation_keys=cfg.write_rules.aggregation_keys)
    wb = excel_reader.open_workbook(data)
    report = build_plan(wb, rows, cfg)

    ready = [i for i in report.items if i.status == Status.READY]
    assert len(ready) == 1
    item = ready[0]
    assert item.sheet == "S1"
    assert item.matched_row == 2          # date 15/05/2026 lives in row 2
    assert item.target_column == column_index_from_string("C")
    assert round(item.final_value, 2) == 13335.00   # 630 + 12,705

    ops = plan_to_write_ops(report, cfg)
    out_bytes, outcomes = apply_writes(data, ops)

    # Verify the written cell and that the formula in E2 survived.
    out_wb = openpyxl.load_workbook(io.BytesIO(out_bytes))
    ws = out_wb["S1"]
    assert round(ws["C2"].value, 2) == 13335.00
    assert ws["E2"].value == "=B2+C2"     # untouched formula preserved
    assert ws["B2"].value is None          # unrelated column untouched

    # Audit must capture the write.
    audit = build_audit_records(outcomes, report.items, source_file="sample.pdf")
    assert len(audit) == 1
    assert audit[0]["target_cell"] == "C2"
    assert round(float(audit[0]["new_value"]), 2) == 13335.00


def test_no_overwrite_without_confirmation():
    data = _make_workbook_bytes()
    # Pre-fill the target cell C2 with an existing value.
    wb = openpyxl.load_workbook(io.BytesIO(data))
    wb["S1"]["C2"] = 999
    buf = io.BytesIO(); wb.save(buf); data = buf.getvalue()

    cfg = _config_for_letter_target()
    cfg.write_rules.write_action = WriteAction.ASK   # default safety: ask/skip

    records = _records()
    rows = aggregate(records, cfg.source, aggregation_keys=cfg.write_rules.aggregation_keys)
    wb2 = excel_reader.open_workbook(data)
    report = build_plan(wb2, rows, cfg)
    # The item must be flagged as a conflict, not silently written.
    assert any(i.status == Status.CONFLICT for i in report.items)

    # Applying with no resolver must skip the conflicting cell (no overwrite).
    ops = plan_to_write_ops(report, cfg)
    out_bytes, outcomes = apply_writes(data, ops)
    out_wb = openpyxl.load_workbook(io.BytesIO(out_bytes))
    assert out_wb["S1"]["C2"].value == 999   # original value untouched


def test_multiple_missing_dates_append_to_distinct_rows():
    """Several missing dates under one group must NOT collide on one row."""
    data = _make_workbook_bytes()  # has dates 15/05 and 16/05 only
    text = """\
Customer Name: SOME GROUP
A1   20/05/2026   100.00
A2   21/05/2026   200.00
A3   22/05/2026   300.00
"""
    records = parse_text_lines([text], group_label="Customer Name",
                               field_names=["Inv No", "Inv Date", "Amount"])
    cfg = _config_for_letter_target()
    cfg.write_rules.insert_missing_date_rows = True
    rows = aggregate(records, cfg.source, aggregation_keys=cfg.write_rules.aggregation_keys)
    wb = excel_reader.open_workbook(data)
    report = build_plan(wb, rows, cfg)

    inserted = [i for i in report.items if i.is_insert]
    assert len(inserted) == 3
    target_rows = [i.matched_row for i in inserted]
    assert len(set(target_rows)) == 3          # distinct rows, no collision

    ops = plan_to_write_ops(report, cfg)
    out_bytes, _ = apply_writes(data, ops)
    out_wb = openpyxl.load_workbook(io.BytesIO(out_bytes))
    ws = out_wb["S1"]
    # Each appended row holds its own amount and its own stamped date.
    by_amount = {}
    for i in inserted:
        by_amount[i.aggregated_amount] = i.matched_row
    assert ws.cell(row=by_amount[100.0], column=3).value == 100.0   # target column "C"
    assert ws.cell(row=by_amount[300.0], column=3).value == 300.0
    # Date stamped into the date column ("A") of each appended row.
    from src import utils
    assert utils.parse_date(ws.cell(row=by_amount[100.0], column=1).value) == dt.date(2026, 5, 20)
    assert ws["E2"].value == "=B2+C2"          # pre-existing formula untouched


def test_backup_is_exact_copy():
    data = _make_workbook_bytes()
    backup = excel_writer.make_backup(data)
    assert backup == data
    assert backup is not data
