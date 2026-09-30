import datetime as dt
import io
from copy import copy
import re
import zipfile

import openpyxl
import pytest

from src.profiles import ColumnRef
from src.workbook import CellChange, Workbook, WorkbookError, update_workbook, verify_update
from tests.conftest import make_workbook, row_of


@pytest.fixture
def book():
    return make_workbook(["S1", "S2"], other_layout=["K1"])


def _parts(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    return {n: z.read(n) for n in z.namelist()}


def test_only_target_values_change_and_formatting_is_kept(book):
    r = row_of(dt.date(2026, 8, 3))
    changes = [CellChange("S1", f"B{r}", 19750.0), CellChange("S2", f"B{r}", 7000)]
    out = update_workbook(book, changes)
    assert verify_update(book, out, changes) == []

    before, after = _parts(book), _parts(out)
    assert set(before) == set(after)
    touched = {n for n in before if before[n] != after[n]}
    assert touched == {"xl/workbook.xml", "xl/worksheets/sheet2.xml", "xl/worksheets/sheet3.xml"}
    assert any("comments" in n or "vmlDrawing" in n for n in before)   # untouched extras exist

    wb_before, wb_after = openpyxl.load_workbook(io.BytesIO(book)), openpyxl.load_workbook(io.BytesIO(out))
    assert wb_after.sheetnames == wb_before.sheetnames
    c0, c1 = wb_before["S1"][f"B{r}"], wb_after["S1"][f"B{r}"]
    assert c1.value == 19750
    style = lambda c: (copy(c.font), copy(c.alignment), c.number_format, copy(c.border), copy(c.fill))
    assert style(c1) == style(c0)
    assert wb_after["S1"][f"D{r}"].value == f"=+B{r}-C{r}"            # formulas preserved
    assert wb_after["Consolidated"]["B2"].value is None
    assert 'fullCalcOnLoad="1"' in after["xl/workbook.xml"].decode()


def test_refuses_to_overwrite_a_formula(book):
    with pytest.raises(WorkbookError):
        update_workbook(book, [CellChange("S1", "D5", 1.0)])


def test_formula_output_and_rewrite_of_own_breakup(book):
    change = [CellChange("S1", "B10", "=12500+7250")]
    out = update_workbook(book, change)
    assert verify_update(book, out, change) == []
    assert openpyxl.load_workbook(io.BytesIO(out))["S1"]["B10"].value == "=12500+7250"
    again = [CellChange("S1", "B10", 19750.0)]                          # own breakup may be replaced
    assert verify_update(out, update_workbook(out, again), again) == []


def test_missing_cell_is_created_in_column_order():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["C1"] = "x", "y"
    buf = io.BytesIO(); wb.save(buf); data = buf.getvalue()
    change = [CellChange("Sheet", "B1", 5)]
    out = update_workbook(data, change)
    assert verify_update(data, out, change) == []
    xml = _parts(out)["xl/worksheets/sheet1.xml"].decode()
    assert re.search(r'r="A1".*r="B1".*r="C1"', xml)


def test_verification_catches_unplanned_changes(book):
    planned = [CellChange("S1", "B10", 1.0)]
    sneaky = update_workbook(book, planned + [CellChange("S1", "B11", 2.0)])
    assert verify_update(book, sneaky, planned)


def test_real_dates_behind_short_display_and_full_column_scan(book):
    wb = Workbook(book)
    index = wb.date_rows("S1", "A", 1)
    late = dt.date(2026, 9, 30)                       # far below any preview window
    assert index[late] == [row_of(late)]
    assert wb.cell("S1", "A2").value.year == 2026     # 'd-mmm' shows no year; value has it
