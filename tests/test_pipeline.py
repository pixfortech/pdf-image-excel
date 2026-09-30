"""End-to-end behaviour: first run -> saved profile -> recognised next time ->
preview -> update the uploaded workbook -> verify."""
import datetime as dt
import io

import openpyxl
import pytest

from src import pipeline, plan
from src.profiles import FORMULA, ColumnRef, SheetLayout
from tests.conftest import BRANCH_LAYOUT, make_report_pdf, make_workbook, row_of

MAPPING = {"ALPHA ONE": "S1", "BETA TWO": "S2", "GAMMA THREE": "K1"}


@pytest.fixture
def files(sample_groups):
    return make_report_pdf(sample_groups), make_workbook(["S1", "S2", "S3"], other_layout=["K1"])


def _configure(store, pdf, xlsx, **write):
    first = pipeline.prepare(pdf, "r.pdf", xlsx, "b.xlsx", store)
    assert first.profile_source == pipeline.NEW
    profile = first.profile
    profile.profile_name = "Group sales"
    profile.group_to_sheet = dict(MAPPING)
    profile.sheet_template = BRANCH_LAYOUT
    for k, v in write.items():
        setattr(profile.write, k, v)
    store.save(profile)
    return profile


def test_recognised_mapping_plans_each_group_into_its_own_sheet(store, files):
    pdf, xlsx = files
    _configure(store, pdf, xlsx)
    p = pipeline.prepare(pdf, "r.pdf", xlsx, "b.xlsx", store)
    assert p.profile_source == pipeline.LOADED and not p.hard_blockers
    assert p.reconciliation.ok

    status = {s.group: s for s in p.plan.sheets}
    assert status["ALPHA ONE"].status == plan.READY_NEW and status["ALPHA ONE"].sheet == "S1"
    assert status["GAMMA THREE"].status == plan.COLUMN_PROBLEM            # K1 differs from template
    assert "no 'CHALLAN' column" in status["GAMMA THREE"].note          # only this sheet is flagged

    day = next(r for r in p.plan.rows if r.total.group == "ALPHA ONE" and r.total.date == dt.date(2026, 8, 3))
    assert (day.sheet, day.cell, day.new_value, day.status) == \
           ("S1", f"B{row_of(dt.date(2026, 8, 3))}", 19750.0, plan.READY_NEW)
    assert not p.can_update() and p.can_update(exclude_blocked=True)


def test_update_writes_only_mapped_cells(store, files):
    pdf, xlsx = files
    _configure(store, pdf, xlsx)
    p = pipeline.prepare(pdf, "r.pdf", xlsx, "b.xlsx", store)
    result = pipeline.update(p, store, exclude_blocked=True)
    before = openpyxl.load_workbook(io.BytesIO(xlsx))
    after = openpyxl.load_workbook(io.BytesIO(result.workbook))

    planned = {(c.sheet, c.cell) for c in result.changes}
    for ws in before.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                new = after[ws.title][cell.coordinate]
                if (ws.title, cell.coordinate) in planned:
                    continue
                assert new.value == cell.value, f"{ws.title}!{cell.coordinate} changed"
    assert {c.sheet for c in result.changes} == {"S1", "S2"}             # not S3, not Consolidated, not K1
    r = row_of(dt.date(2026, 8, 3))
    assert after["S1"][f"B{r}"].value == 19750 and after["S2"][f"B{r}"].value == 1000
    assert after["S1"][f"C{r}"].value is None                            # blank returns never written
    assert store.load("Group sales").last_success                        # remembered after success

    # Running again on the updated workbook: everything is already there.
    again = pipeline.prepare(pdf, "r.pdf", result.workbook, "b.xlsx", store)
    assert again.ready_count == 0
    with pytest.raises(pipeline.UpdateError, match="Nothing is ready to write"):
        pipeline.update(again, store, exclude_blocked=True)


def test_existing_return_values_survive_and_zero_fill_is_opt_in(store, files):
    pdf, _ = files
    r = row_of(dt.date(2026, 8, 3))
    xlsx = make_workbook(["S1", "S2"], other_layout=["K1"], prefill={f"S1!C{r}": 250})
    _configure(store, pdf, xlsx)
    p = pipeline.prepare(pdf, "r.pdf", xlsx, "b.xlsx", store)
    assert not any(row.field == "returns" for row in p.plan.rows)
    out = pipeline.update(p, store, exclude_blocked=True).workbook
    assert openpyxl.load_workbook(io.BytesIO(out))["S1"][f"C{r}"].value == 250

    profile = store.load("Group sales")
    profile.write.zero_fill_returns = True
    store.save(profile)
    p = pipeline.prepare(pdf, "r.pdf", xlsx, "b.xlsx", store)
    returns = [row for row in p.plan.rows if row.field == "returns" and row.sheet == "S1"]
    assert returns and all(row.new_value == 0.0 for row in returns)


def test_override_resolves_a_differently_laid_out_sheet(store, files):
    pdf, xlsx = files
    profile = _configure(store, pdf, xlsx)
    profile.sheet_overrides["K1"] = SheetLayout(1, ColumnRef("A", "DATE"), ColumnRef("B", "KARKHANA"),
                                                ColumnRef("D", "RETURN"))
    store.save(profile)
    p = pipeline.prepare(pdf, "r.pdf", xlsx, "b.xlsx", store)
    assert not p.plan.blocking and p.can_update()
    gamma = next(r for r in p.plan.rows if r.total.group == "GAMMA THREE")
    assert gamma.sheet == "K1" and gamma.cell.startswith("B") and gamma.new_value == 9999.0


def test_unmapped_group_blocks_until_mapped(store, files):
    pdf, xlsx = files
    profile = _configure(store, pdf, xlsx)
    del profile.group_to_sheet["BETA TWO"]
    store.save(profile)
    p = pipeline.prepare(pdf, "r.pdf", xlsx, "b.xlsx", store)
    assert any(s.status == plan.UNMAPPED for s in p.plan.sheets)
    assert not p.can_update()


def test_totals_that_do_not_reconcile_block_the_update(store, sample_groups):
    xlsx = make_workbook(["S1", "S2"], other_layout=["K1"])
    good = make_report_pdf(sample_groups)
    _configure(store, good, xlsx)
    bad = make_report_pdf(sample_groups, grand_total="1.00")
    p = pipeline.prepare(bad, "r.pdf", xlsx, "b.xlsx", store)
    assert any("do not reconcile" in b for b in p.hard_blockers)
    assert not p.can_update(exclude_blocked=True)


def test_formula_mode_writes_breakup_and_is_idempotent(store, files):
    pdf, xlsx = files
    _configure(store, pdf, xlsx, output_mode=FORMULA)
    p = pipeline.prepare(pdf, "r.pdf", xlsx, "b.xlsx", store)
    out = pipeline.update(p, store, exclude_blocked=True).workbook
    r = row_of(dt.date(2026, 8, 3))
    assert openpyxl.load_workbook(io.BytesIO(out))["S1"][f"B{r}"].value == "=12500+7250"
    again = pipeline.prepare(pdf, "r.pdf", out, "b.xlsx", store)
    assert again.ready_count == 0
