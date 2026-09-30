"""Builders for realistic test files (neutral names and numbers only).

``make_report_pdf`` draws a genuine PDF laid out like a grouped sales report:
title, From/To line, "Group Name: X" sections, a (optionally ruled) header,
right-aligned amounts, group totals, a grand total and page footers, with
sections continuing across pages.  ``make_workbook`` builds branch sheets with
styled cells, formulas, a cell comment and a consolidated sheet.
"""
from __future__ import annotations

import datetime as dt
import io
from typing import Dict, List, Optional, Sequence, Tuple

import openpyxl
import pymupdf
import pytest
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font

from src.profiles import ColumnRef, ProfileStore, SheetLayout

Row = Tuple[str, str, Optional[str], Optional[str]]   # ref, date, amount, returns

HEADER = ["Ref No", "Doc Date", "Total Amount", "Returns"]
CELLS = [(57, 136), (136, 205), (205, 288), (288, 381)]
FONT, SIZE = "helv", 9


def _text(page, x, y, text, right_edge=None):
    if right_edge is not None:
        x = right_edge - pymupdf.get_text_length(text, fontname=FONT, fontsize=SIZE)
    page.insert_text((x, y), text, fontname=FONT, fontsize=SIZE)


def make_report_pdf(groups: Dict[str, List[Row]], *, ruled: bool = True, rows_per_page: int = 30,
                    group_label: str = "Group Name", period=("01/08/2026", "31/08/2026"),
                    group_totals: Optional[Dict[str, Sequence[str]]] = None,
                    grand_total: Optional[str] = "auto") -> bytes:
    doc = pymupdf.open()
    state = {"page": None, "y": 0.0, "lines": 0}

    def new_page(with_header: bool):
        state["page"] = doc.new_page(width=652, height=919)
        state["y"], state["lines"] = 40.0, 0
        if with_header:
            header()

    def advance():
        state["y"] += 16
        state["lines"] += 1

    def header():
        top = state["y"] - 12
        if ruled:
            for x0, x1 in CELLS:
                state["page"].draw_rect(pymupdf.Rect(x0, top, x1, top + 16), color=(0, 0, 0), width=0.6)
        for (x0, x1), label in zip(CELLS, HEADER):
            _text(state["page"], (x0 + x1) / 2 - pymupdf.get_text_length(label, fontname=FONT, fontsize=SIZE) / 2,
                  state["y"], label)
        advance()

    def line(parts):
        if state["lines"] >= rows_per_page:
            new_page(with_header=True)
        for x, text, right in parts:
            _text(state["page"], x, state["y"], text, right_edge=right)
        advance()

    new_page(with_header=False)
    _text(state["page"], 160, state["y"], "Group wise sales details"); advance()
    _text(state["page"], 72, state["y"], f"From Date: {period[0]} ToDate: {period[1]}"); advance()

    overall = 0.0
    for name, rows in groups.items():
        line([(75, f"{group_label}: {name}", None)])
        if state["lines"] >= rows_per_page:
            new_page(with_header=False)
        header()
        amount_sum = returns_sum = 0.0
        for ref, date, amount, returns in rows:
            parts = [(77, ref, None), (147, date, None)]
            if amount:
                parts.append((0, amount, 286))
                amount_sum += float(amount.replace(",", ""))
            if returns:
                parts.append((0, returns, 378))
                returns_sum += float(returns.replace(",", ""))
            line(parts)
        overall += amount_sum
        printed = (group_totals or {}).get(name)
        if printed is None:
            printed = [f"{amount_sum:,.2f}"] + ([f"{returns_sum:,.2f}"] if returns_sum else [])
        line([(75, f"{group_label.split()[0]} Total : " + " ".join(printed), None)])
    if grand_total:
        line([(75, f"Total : {overall:,.2f}" if grand_total == "auto" else f"Total : {grand_total}", None)])

    for i, page in enumerate(doc, start=1):
        _text(page, 72, 900, f"Total NoOfPages: {doc.page_count} Page No: {i}")
    return doc.tobytes()


BRANCH = ["DATE", "CHALLAN", "RETURN", "TOTAL"]
OTHER = ["DATE", "KARKHANA", "BD", "RETURN", "TOTAL"]


def make_workbook(branch_sheets: Sequence[str], *, other_layout: Sequence[str] = (),
                  start=dt.date(2026, 7, 1), days=92, prefill: Optional[Dict[str, object]] = None) -> bytes:
    """Branch sheets: DATE | CHALLAN | RETURN | TOTAL(=B-C), styled like the
    real file.  ``other_layout`` sheets use a different column pattern."""
    wb = openpyxl.Workbook()
    cons = wb.active
    cons.title = "Consolidated"
    cons.append(["DATE"] + list(branch_sheets))
    for sheet_names, header in ((branch_sheets, BRANCH), (other_layout, OTHER)):
        for name in sheet_names:
            ws = wb.create_sheet(name)
            ws.append(header)
            for i in range(days):
                r = i + 2
                ws.cell(row=r, column=1, value=dt.datetime.combine(start + dt.timedelta(days=i), dt.time()))
                ws.cell(row=r, column=1).number_format = "d-mmm"
                total_col = header.index("TOTAL") + 1
                ws.cell(row=r, column=total_col,
                        value=f"=+B{r}-C{r}" if header is BRANCH else f"=+(B{r}+C{r})-D{r}")
                for c in range(1, len(header) + 1):
                    cell = ws.cell(row=r, column=c)
                    cell.font = Font(name="Tenorite", size=11)
                    cell.alignment = Alignment(horizontal="center")
            ws["A1"].comment = Comment("Header note", "tester")
            ws.column_dimensions["B"].width = 14
    for i in range(days):
        cons.cell(row=i + 2, column=1, value=dt.datetime.combine(start + dt.timedelta(days=i), dt.time()))
    for ref, value in (prefill or {}).items():
        sheet, cell = ref.split("!")
        wb[sheet][cell] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def row_of(date: dt.date, start=dt.date(2026, 7, 1)) -> int:
    return (date - start).days + 2


BRANCH_LAYOUT = SheetLayout(1, ColumnRef("A", "DATE"), ColumnRef("B", "CHALLAN"), ColumnRef("C", "RETURN"))


@pytest.fixture
def store(tmp_path) -> ProfileStore:
    return ProfileStore(tmp_path / "profiles")


@pytest.fixture
def sample_groups() -> Dict[str, List[Row]]:
    """Group A spans a page break and has two invoices on one date."""
    a = [("REF-1001", "03/08/2026", "12,500.00", None),
         ("REF-1002", "03/08/2026", "7,250.00", None),
         ("REF-1003", "04/08/2026", "4,100.00", None),
         ("REF-1004", "04/08/2026", "2,900.00", None)]
    a += [(f"REF-11{i:02d}", f"{5 + i // 2:02d}/08/2026", f"{100 + i:,.2f}", None) for i in range(40)]
    return {
        "ALPHA ONE": a,
        "BETA TWO": [("REF-2001", "03/08/2026", "1,000.00", None),
                     ("REF-2002", "05/08/2026", "2,500.00", None)],
        "GAMMA THREE": [("REF-3001", "06/08/2026", "9,999.00", None)],
    }
