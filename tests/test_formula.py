"""Tests for the numeric-total vs Excel-formula-breakup behaviour.

Confirms (with arbitrary, non-hardcoded values supplied at test time):
* duplicate group/date rows sum to a clean numeric total,
* the formula export starts with '=' and contains no thousands separators,
* numeric write mode writes the numeric total into the cell,
* formula write mode writes the '=...' formula into the cell,
* the numeric total is always kept alongside the formula.
"""
import io

import openpyxl

from src import excel_reader
from src.aggregator import aggregate, build_formula, grouped_export_rows
from src.excel_writer import apply_writes
from src.mapping import (
    AppConfig, OutputType, SheetMapping, SourceMapping, TargetMode,
    WriteAction, WriteRules,
)
from src.parser import parse_text_lines
from src.validator import Status, build_plan, plan_to_export_rows, plan_to_write_ops


def _records():
    # Two invoices on the SAME date under the same group: 630.00 and 12,705.00.
    text = """\
Customer Name: SOME GROUP
A1   15/05/2026   630.00
A2   15/05/2026   12,705.00
"""
    return parse_text_lines([text], group_label="Customer Name",
                            field_names=["Inv No", "Inv Date", "Amount"])


def _source():
    return SourceMapping(date_field="Inv Date", amount_field="Amount",
                         sum_duplicate_dates=True, thousands_sep=",", decimal_sep=".")


def test_build_formula_strips_commas_and_leads_with_equals():
    assert build_formula([630.0, 12705.0]) == "=630+12705"
    # No thousands separators ever appear.
    assert "," not in build_formula([1234567.0, 1000.0])
    assert build_formula([1234567.0, 1000.0]).startswith("=")
    # Negative components render as subtraction, not "+-".
    assert build_formula([100.0, -20.0]) == "=100-20"
    # Decimals kept only when needed.
    assert build_formula([630.5, 12705.0]) == "=630.5+12705"
    assert build_formula([]) == "=0"


def test_numeric_total_and_formula_breakup_both_present():
    rows = aggregate(_records(), _source(), aggregation_keys=["group", "date"])
    assert len(rows) == 1
    row = rows[0]
    # 1) numeric total is clean
    assert round(row.amount, 2) == 13335.00
    # 2) human-readable breakup keeps the original comma formatting
    assert row.invoice_breakup == "630.00 + 12,705.00"
    # 3) formula breakup is comma-free and starts with '='
    assert row.excel_formula_breakup == "=630+12705"


def test_grouped_export_has_three_value_columns():
    rows = aggregate(_records(), _source(), aggregation_keys=["group", "date"])
    export = grouped_export_rows(rows)[0]
    assert export["aggregated_amount"] == 13335.00            # numeric
    assert export["invoice_breakup"] == "630.00 + 12,705.00"  # preview
    assert export["excel_formula_breakup"] == "=630+12705"    # formula
    assert "," not in export["excel_formula_breakup"]


def _workbook_bytes():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "S1"
    ws.append(["Date", "Amount"])
    import datetime as dt
    ws.append([dt.date(2026, 5, 15), None])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _config(output_type):
    return AppConfig(
        source=_source(),
        group_to_sheet={"SOME GROUP": "S1"},
        sheets={"S1": SheetMapping(
            sheet_name="S1", header_row=1,
            date_target_mode=TargetMode.COLUMN_LETTER, date_column="A",
            amount_target_mode=TargetMode.COLUMN_LETTER, amount_column="B",
        )},
        write_rules=WriteRules(write_action=WriteAction.REPLACE,
                               output_type=output_type, aggregation_keys=["group", "date"]),
    )


def _run(output_type):
    data = _workbook_bytes()
    cfg = _config(output_type)
    rows = aggregate(_records(), cfg.source, aggregation_keys=["group", "date"])
    report = build_plan(excel_reader.open_workbook(data), rows, cfg)
    ops = plan_to_write_ops(report, cfg)
    out, outcomes = apply_writes(data, ops)
    return cfg, report, out


def test_numeric_write_mode_writes_numeric_total():
    cfg, report, out = _run(OutputType.NUMERIC)
    item = [i for i in report.items if i.status == Status.READY][0]
    assert round(item.aggregated_amount, 2) == 13335.00
    assert item.final_value == 13335.0          # numeric, not a formula
    ws = openpyxl.load_workbook(io.BytesIO(out))["S1"]
    assert ws["B2"].value == 13335.0
    assert ws["B2"].data_type != "f"            # stored as a number


def test_formula_write_mode_writes_formula():
    cfg, report, out = _run(OutputType.FORMULA)
    item = [i for i in report.items if i.status == Status.READY][0]
    # Numeric total is STILL kept on the plan item alongside the formula.
    assert round(item.aggregated_amount, 2) == 13335.00
    assert item.excel_formula_breakup == "=630+12705"
    assert item.final_value == "=630+12705"
    ws = openpyxl.load_workbook(io.BytesIO(out))["S1"]
    assert ws["B2"].value == "=630+12705"
    assert ws["B2"].data_type == "f"            # stored as a real formula


def test_write_plan_export_contains_all_three_values():
    cfg, report, _ = _run(OutputType.NUMERIC)
    plan = plan_to_export_rows(report, cfg)
    ready = [r for r in plan if r["status"] == Status.READY][0]
    assert ready["aggregated_amount"] == 13335.00
    assert ready["invoice_breakup"] == "630.00 + 12,705.00"
    assert ready["excel_formula_breakup"] == "=630+12705"
    assert ready["write_mode"] == "numeric"
    assert ready["value_to_write"] == 13335.0
