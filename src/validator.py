"""Validation and build of the final write plan.

Combines aggregated source rows with the Excel mapping to produce a list of
:class:`PlanItem` rows, each annotated with a status (Ready, Unmapped Group,
Missing Column, Date Not Found, Conflict, Error).  It also resolves the actual
:class:`~src.excel_writer.WriteOp` objects for the items that are ready.

No worksheet, column, customer or value is hardcoded; everything comes from the
:class:`~src.mapping.AppConfig` provided by the user.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Dict, List, Optional

from . import excel_reader
from .aggregator import AggregatedRow, build_formula
from .excel_writer import WriteOp
from .mapping import (
    AppConfig,
    DateNotFoundAction,
    OutputType,
    TargetMode,
    WriteAction,
)
from . import utils


class Status:
    READY = "Ready"
    UNMAPPED_GROUP = "Unmapped Group"
    MISSING_COLUMN = "Missing Column"
    SHEET_NOT_FOUND = "Sheet Not Found"
    DATE_NOT_FOUND = "Date Not Found"
    CONFLICT = "Conflict"
    ERROR = "Error"


@dataclass
class PlanItem:
    group: str
    sheet: str = ""
    date: object = None
    date_raw: str = ""
    matched_row: Optional[int] = None
    target_column: Optional[int] = None
    target_column_letter: str = ""
    target_cell: str = ""
    source_amounts: List[float] = field(default_factory=list)
    aggregated_amount: float = 0.0          # numeric total, ALWAYS kept
    invoice_breakup: str = ""               # human-readable, e.g. "630.00 + 12,705.00"
    excel_formula_breakup: str = ""         # formula, e.g. "=630+12705"
    return_amount: float = 0.0
    existing_value: object = None
    final_value: object = None              # numeric total or formula (per write mode)
    write_action: str = ""
    status: str = Status.READY
    messages: List[str] = field(default_factory=list)
    return_column: Optional[int] = None
    return_existing_value: object = None
    source_records: List[dict] = field(default_factory=list)
    is_insert: bool = False
    is_append: bool = False          # appended below existing data (no physical row insert)
    template_row: Optional[int] = None
    date_column: Optional[int] = None  # resolved date column index (for inserts)


@dataclass
class ValidationReport:
    items: List[PlanItem] = field(default_factory=list)
    unmapped_groups: List[str] = field(default_factory=list)
    dates_not_found: int = 0
    missing_columns: int = 0
    conflicts: int = 0
    errors: int = 0
    duplicate_headers: Dict[str, List[str]] = field(default_factory=dict)

    @property
    def has_blocking_errors(self) -> bool:
        return (
            bool(self.unmapped_groups)
            or self.missing_columns > 0
            or self.errors > 0
        )

    @property
    def ready_count(self) -> int:
        return sum(1 for i in self.items if i.status == Status.READY)


def build_plan(
    wb,
    rows: List[AggregatedRow],
    config: AppConfig,
) -> ValidationReport:
    report = ValidationReport()
    sheet_names = set(excel_reader.list_sheets(wb))

    # Pre-compute duplicate header warnings per mapped sheet.
    for sheet_name, sm in config.sheets.items():
        if sheet_name in sheet_names:
            dups = excel_reader.find_duplicate_headers(wb, sheet_name, sm.header_row)
            if dups:
                report.duplicate_headers[sheet_name] = dups

    # Per-sheet cursor so that multiple missing dates appended to the same sheet
    # land on distinct, sequential rows instead of all colliding at max_row+1.
    append_cursor: Dict[str, int] = {}

    for agg in rows:
        item = PlanItem(
            group=agg.group,
            date=agg.date,
            date_raw=agg.date_raw,
            source_amounts=[r.get(config.source.amount_field, "") for r in agg.source_records]
            if config.source.amount_field else [],
            aggregated_amount=agg.amount,
            invoice_breakup=agg.invoice_breakup,
            excel_formula_breakup=agg.excel_formula_breakup,
            return_amount=agg.return_amount,
            source_records=agg.source_records,
        )
        item.messages.extend(agg.warnings)

        sheet = config.group_to_sheet.get(agg.group, "")
        if not sheet:
            item.status = Status.UNMAPPED_GROUP
            item.messages.append("No worksheet mapped for this group.")
            report.items.append(item)
            if agg.group not in report.unmapped_groups:
                report.unmapped_groups.append(agg.group)
            continue
        item.sheet = sheet

        if sheet not in sheet_names:
            item.status = Status.SHEET_NOT_FOUND
            item.messages.append(f"Worksheet '{sheet}' not found in workbook.")
            report.errors += 1
            report.items.append(item)
            continue

        sm = config.sheets.get(sheet)
        if sm is None:
            item.status = Status.MISSING_COLUMN
            item.messages.append(f"No Excel column mapping configured for '{sheet}'.")
            report.missing_columns += 1
            report.items.append(item)
            continue

        # Resolve amount target.
        if sm.amount_target_mode == TargetMode.CELL_REFERENCE and sm.amount_cell:
            rc = excel_reader.cell_reference_to_rc(sm.amount_cell)
            if rc is None:
                item.status = Status.ERROR
                item.messages.append(f"Invalid target cell: {sm.amount_cell}")
                report.errors += 1
                report.items.append(item)
                continue
            item.matched_row, item.target_column = rc
            item.target_cell = sm.amount_cell.upper()
            item.target_column_letter = "".join(ch for ch in item.target_cell if ch.isalpha())
        else:
            col = excel_reader.resolve_column_index(
                wb, sheet,
                mode=sm.amount_target_mode.value,
                selector=sm.amount_column,
                header_row=sm.header_row,
            )
            if col is None:
                item.status = Status.MISSING_COLUMN
                item.messages.append(
                    f"Target amount column '{sm.amount_column}' not found in '{sheet}'."
                )
                report.missing_columns += 1
                report.items.append(item)
                continue
            item.target_column = col
            from openpyxl.utils import get_column_letter
            item.target_column_letter = get_column_letter(col)

            # Resolve the matching row by date.
            date_col = excel_reader.resolve_column_index(
                wb, sheet,
                mode=sm.date_target_mode.value,
                selector=sm.date_column,
                header_row=sm.header_row,
            )
            if date_col is None and sm.date_column:
                item.status = Status.MISSING_COLUMN
                item.messages.append(f"Date column '{sm.date_column}' not found in '{sheet}'.")
                report.missing_columns += 1
                report.items.append(item)
                continue

            matched = excel_reader.find_row_by_date(
                wb, sheet, date_col, agg.date,
                header_row=sm.header_row,
                date_formats=sm.date_formats or None,
            ) if date_col else None

            item.date_column = date_col
            if matched is None:
                action = config.write_rules.date_not_found_action
                do_insert = (action in (DateNotFoundAction.INSERT_ROW, DateNotFoundAction.COPY_NEAREST)
                             or config.write_rules.insert_missing_date_rows)
                if do_insert:
                    ws = wb[sheet]
                    base = _insert_position(wb, sheet, date_col, agg.date, sm)
                    cursor = append_cursor.get(sheet)
                    if base > ws.max_row:
                        # Append below existing data: distinct sequential rows.
                        row = cursor if cursor is not None else base
                        item.is_append = True
                        item.messages.append("Date not found; will append a new row.")
                    else:
                        # Mid-sheet insert: offset by inserts already planned above.
                        row = base if cursor is None else max(base, cursor)
                        item.messages.append("Date not found; will insert a new row.")
                    append_cursor[sheet] = row + 1
                    item.is_insert = True
                    item.matched_row = row
                    item.template_row = excel_reader.nearest_date_row(
                        wb, sheet, date_col, agg.date,
                        header_row=sm.header_row, date_formats=sm.date_formats or None,
                    )
                else:
                    item.status = Status.DATE_NOT_FOUND
                    item.messages.append("Matching date row not found in worksheet.")
                    report.dates_not_found += 1
                    report.items.append(item)
                    continue
            else:
                item.matched_row = matched
            item.target_cell = f"{item.target_column_letter}{item.matched_row}"

        # Resolve optional return column.
        if config.source.return_field and (sm.return_column or sm.return_cell):
            if sm.return_target_mode == TargetMode.CELL_REFERENCE and sm.return_cell:
                rc = excel_reader.cell_reference_to_rc(sm.return_cell)
                if rc:
                    item.return_column = rc[1]
            else:
                item.return_column = excel_reader.resolve_column_index(
                    wb, sheet,
                    mode=sm.return_target_mode.value,
                    selector=sm.return_column,
                    header_row=sm.header_row,
                )

        # Compute existing value & final value.
        if not item.is_insert and item.matched_row and item.target_column:
            ws = wb[sheet]
            item.existing_value = ws.cell(row=item.matched_row, column=item.target_column).value
            if item.return_column:
                item.return_existing_value = ws.cell(row=item.matched_row, column=item.return_column).value

        item.final_value = _compute_final_value(item, agg, config)
        item.write_action = config.write_rules.write_action.value

        # Conflict detection.
        if (
            not item.is_insert
            and item.existing_value not in (None, "")
            and config.write_rules.write_action in (WriteAction.ASK, WriteAction.SKIP_IF_VALUE)
        ):
            item.status = Status.CONFLICT
            item.messages.append("Target cell already has a value.")
            report.conflicts += 1
        else:
            item.status = Status.READY

        report.items.append(item)

    return report


def _insert_position(wb, sheet, date_col, target_date, sm) -> int:
    """Choose a row index to insert at: keep dates sorted when possible."""
    ws = wb[sheet]
    if not date_col or target_date is None:
        return ws.max_row + 1
    insert_at = ws.max_row + 1
    for r in range(sm.header_row + 1, ws.max_row + 1):
        val = ws.cell(row=r, column=date_col).value
        parsed = utils.parse_date(val, formats=sm.date_formats or None)
        if parsed is not None and parsed > target_date:
            insert_at = r
            break
    return insert_at


def _compute_final_value(item: PlanItem, agg: AggregatedRow, config: AppConfig):
    """Pick the value to write based on the chosen output mode.

    The numeric total (``item.aggregated_amount``) is always preserved on the
    plan item; this only decides what lands in the target cell.
    """
    if config.write_rules.output_type == OutputType.FORMULA:
        # Use the precomputed comma-free formula, e.g. "=630+12705".
        return agg.excel_formula_breakup or build_formula([agg.amount])
    if config.write_rules.write_action == WriteAction.ADD:
        base = item.existing_value if isinstance(item.existing_value, (int, float)) else 0
        return round((base or 0) + agg.amount, 4)
    return round(agg.amount, 4)


def plan_to_write_ops(report: ValidationReport, config: AppConfig) -> List[WriteOp]:
    """Convert READY (and conflict-resolved) plan items into write operations."""
    ops: List[WriteOp] = []
    for item in report.items:
        if item.status not in (Status.READY, Status.CONFLICT):
            continue
        if item.target_column is None or item.matched_row is None:
            continue
        comment = ""
        if config.write_rules.add_source_comment and item.source_records:
            lines = [str(r.get("_source_text", "")) for r in item.source_records]
            comment = "Source rows:\n" + "\n".join(l for l in lines if l)

        # For inserted/appended rows, also stamp the date into the date column so
        # the new row is identifiable. Whichever op for this row comes first
        # performs the physical insert / formatting copy.
        date_op_added = False
        if item.is_insert and item.date_column and item.date is not None:
            ops.append(WriteOp(
                sheet_name=item.sheet,
                row=item.matched_row,
                column=item.date_column,
                value=item.date,
                action=WriteAction.REPLACE,
                output_type=OutputType.NUMERIC,
                insert_row=item.is_insert and not item.is_append,
                append_only=item.is_append,
                template_row=item.template_row,
                label=f"{item.group} / {item.date_raw} (date)",
            ))
            date_op_added = True

        ops.append(WriteOp(
            sheet_name=item.sheet,
            row=item.matched_row,
            column=item.target_column,
            value=item.final_value,
            action=config.write_rules.write_action,
            output_type=config.write_rules.output_type,
            comment=comment,
            # If a date op already inserted/formatted this row, the amount op just
            # writes its cell; otherwise the amount op performs the insert/append.
            insert_row=item.is_insert and not item.is_append and not date_op_added,
            append_only=item.is_append and not date_op_added,
            template_row=item.template_row,
            label=f"{item.group} / {item.date_raw}",
        ))

        # Optional return write.
        if item.return_column and config.source.return_field:
            ops.append(WriteOp(
                sheet_name=item.sheet,
                row=item.matched_row,
                column=item.return_column,
                value=round(item.return_amount, 4),
                action=config.write_rules.write_action,
                output_type=OutputType.NUMERIC,
                label=f"{item.group} / {item.date_raw} (return)",
            ))
    return ops


def plan_to_export_rows(report: ValidationReport, config: AppConfig) -> List[dict]:
    """Flatten the write plan into export rows (for write_plan.csv).

    Always includes the numeric ``aggregated_amount``, the human-readable
    ``invoice_breakup`` and the ``excel_formula_breakup`` so all three value
    forms are visible regardless of the chosen write mode.
    """
    mode = config.write_rules.output_type.value
    out: List[dict] = []
    for it in report.items:
        out.append({
            "customer_name": it.group,
            "worksheet": it.sheet,
            "date": it.date_raw,
            "num_source_rows": len(it.source_records),
            "invoice_breakup": it.invoice_breakup,
            "aggregated_amount": round(it.aggregated_amount, 2),
            "excel_formula_breakup": it.excel_formula_breakup,
            "return_amount": round(it.return_amount, 2),
            "matched_row": it.matched_row,
            "target_column": it.target_column_letter,
            "target_cell": it.target_cell,
            "existing_value": it.existing_value,
            "write_mode": mode,
            "value_to_write": it.final_value,
            "status": it.status,
            "messages": "; ".join(it.messages),
        })
    return out
