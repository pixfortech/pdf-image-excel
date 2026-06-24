"""Reading existing Excel workbooks with openpyxl.

Provides helpers to list sheets, preview cells, resolve a target column by
header name / column letter / cell reference, and locate a row by matching a
date in a chosen date column.  Repeated header names are handled by always
allowing column-letter and cell-reference selection in addition to header name.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from typing import List, Optional, Sequence, Tuple

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string

from . import utils


@dataclass
class CellPreview:
    coordinate: str
    value: object


def open_workbook(source, *, data_only: bool = False):
    """Open a workbook from a path, bytes, or file-like object.

    ``data_only=False`` preserves formulas (the default for writing).  We never
    modify the object that callers pass in beyond what openpyxl needs.
    """
    if isinstance(source, (bytes, bytearray)):
        return load_workbook(io.BytesIO(bytes(source)), data_only=data_only)
    if hasattr(source, "read"):
        return load_workbook(io.BytesIO(source.read()), data_only=data_only)
    return load_workbook(source, data_only=data_only)


def list_sheets(wb) -> List[str]:
    return list(wb.sheetnames)


def preview_sheet(wb, sheet_name: str, max_rows: int = 30, max_cols: int = 30) -> List[List[object]]:
    ws = wb[sheet_name]
    rows: List[List[object]] = []
    for r in range(1, min(ws.max_row, max_rows) + 1):
        row: List[object] = []
        for c in range(1, min(ws.max_column, max_cols) + 1):
            row.append(ws.cell(row=r, column=c).value)
        rows.append(row)
    return rows


def header_values(wb, sheet_name: str, header_row: int) -> List[Tuple[str, str]]:
    """Return ``(column_letter, header_text)`` pairs for a header row."""
    ws = wb[sheet_name]
    out: List[Tuple[str, str]] = []
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=c).value
        out.append((get_column_letter(c), "" if val is None else str(val)))
    return out


def find_duplicate_headers(wb, sheet_name: str, header_row: int) -> List[str]:
    seen = {}
    dups = []
    for _, text in header_values(wb, sheet_name, header_row):
        key = utils.normalise_whitespace(text).lower()
        if not key:
            continue
        seen[key] = seen.get(key, 0) + 1
        if seen[key] == 2:
            dups.append(text)
    return dups


def resolve_column_index(
    wb,
    sheet_name: str,
    *,
    mode: str,
    selector: str,
    header_row: int = 1,
) -> Optional[int]:
    """Resolve a 1-based column index from a selector.

    ``mode`` is one of ``header_name``, ``column_letter``, ``cell_reference``.
    For ``cell_reference`` the column part of the reference is used.
    Returns ``None`` when it cannot be resolved.
    """
    if not selector:
        return None
    if mode == "column_letter":
        try:
            return column_index_from_string(selector.strip().upper())
        except Exception:
            return None
    if mode == "cell_reference":
        try:
            col_letter, _row = coordinate_from_string(selector.strip().upper())
            return column_index_from_string(col_letter)
        except Exception:
            return None
    # header_name (first match)
    ws = wb[sheet_name]
    target = utils.normalise_whitespace(selector).lower()
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=c).value
        if val is not None and utils.normalise_whitespace(str(val)).lower() == target:
            return c
    return None


def cell_reference_to_rc(ref: str) -> Optional[Tuple[int, int]]:
    try:
        col_letter, row = coordinate_from_string(ref.strip().upper())
        return row, column_index_from_string(col_letter)
    except Exception:
        return None


def find_row_by_date(
    wb,
    sheet_name: str,
    date_col_index: int,
    target_date,
    *,
    header_row: int = 1,
    date_formats: Optional[Sequence[str]] = None,
    interpretation: str = "dmy",
) -> Optional[int]:
    """Find the 1-based row whose date column matches ``target_date``.

    Handles real date objects, datetimes, Excel serial numbers and text dates in
    many formats; everything is normalised to a bare ``date`` before comparing.
    Returns ``None`` if no row matches.
    """
    if target_date is None or date_col_index is None:
        return None
    ws = wb[sheet_name]
    for r in range(header_row + 1, ws.max_row + 1):
        val = ws.cell(row=r, column=date_col_index).value
        parsed = utils.parse_date(val, formats=date_formats, interpretation=interpretation)
        if parsed is not None and parsed == target_date:
            return r
    return None


def nearest_date_row(
    wb,
    sheet_name: str,
    date_col_index: int,
    target_date,
    *,
    header_row: int = 1,
    date_formats: Optional[Sequence[str]] = None,
    interpretation: str = "dmy",
) -> Optional[int]:
    """Return the row whose date is closest to ``target_date`` (for templating)."""
    if target_date is None or date_col_index is None:
        return None
    ws = wb[sheet_name]
    best_row = None
    best_delta = None
    for r in range(header_row + 1, ws.max_row + 1):
        val = ws.cell(row=r, column=date_col_index).value
        parsed = utils.parse_date(val, formats=date_formats, interpretation=interpretation)
        if parsed is None:
            continue
        delta = abs((parsed - target_date).days)
        if best_delta is None or delta < best_delta:
            best_delta = delta
            best_row = r
    return best_row


def column_dates(
    wb,
    sheet_name: str,
    col_index: int,
    *,
    header_row: int = 1,
    date_formats: Optional[Sequence[str]] = None,
    interpretation: str = "dmy",
    limit: Optional[int] = None,
):
    """Return ``(row, raw_value, parsed_date)`` for cells in a column that parse
    as dates.  Used for previews and date-column suggestions."""
    ws = wb[sheet_name]
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        val = ws.cell(row=r, column=col_index).value
        parsed = utils.parse_date(val, formats=date_formats, interpretation=interpretation)
        if parsed is not None:
            out.append((r, val, parsed))
            if limit and len(out) >= limit:
                break
    return out


def analyze_date_column(
    wb,
    sheet_name: str,
    col_index: int,
    target_dates,
    *,
    header_row: int = 1,
    date_formats: Optional[Sequence[str]] = None,
    interpretation: str = "dmy",
) -> dict:
    """Summarise how well a column's dates cover ``target_dates`` (a set of dates).

    Returns matched/missing counts, the parseable-date count, and the first few
    parsed dates for display.
    """
    targets = set(d for d in target_dates if d is not None)
    found = set()
    parsed_count = 0
    sample = []
    ws = wb[sheet_name]
    for r in range(header_row + 1, ws.max_row + 1):
        val = ws.cell(row=r, column=col_index).value
        parsed = utils.parse_date(val, formats=date_formats, interpretation=interpretation)
        if parsed is None:
            continue
        parsed_count += 1
        if len(sample) < 10:
            sample.append((val, parsed))
        if parsed in targets:
            found.add(parsed)
    matched = len(found)
    missing = len(targets) - matched
    return {
        "col_index": col_index,
        "parsed_count": parsed_count,
        "matched": matched,
        "missing": missing,
        "total_targets": len(targets),
        "sample": sample,
    }


def suggest_date_columns(
    wb,
    sheet_name: str,
    target_dates,
    *,
    header_row: int = 1,
    date_formats: Optional[Sequence[str]] = None,
    interpretation: str = "dmy",
    top: int = 3,
):
    """Scan every column and return those that best match ``target_dates``.

    Returns a list of ``(column_letter, header_text, matched_count, parsed_count)``
    sorted by match count (then by parseable-date count), best first.
    """
    ws = wb[sheet_name]
    results = []
    for c in range(1, ws.max_column + 1):
        info = analyze_date_column(
            wb, sheet_name, c, target_dates,
            header_row=header_row, date_formats=date_formats, interpretation=interpretation)
        if info["parsed_count"] == 0:
            continue
        header = ws.cell(row=header_row, column=c).value
        results.append((
            get_column_letter(c),
            "" if header is None else str(header),
            info["matched"],
            info["parsed_count"],
        ))
    results.sort(key=lambda t: (t[2], t[3]), reverse=True)
    return results[:top]
