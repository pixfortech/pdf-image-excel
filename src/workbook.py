"""Read the uploaded workbook, and update it in place.

Reading uses openpyxl (real stored values: a ``d-mmm`` cell showing ``04-Jan``
is read as its full date).

Writing does NOT round-trip the workbook through openpyxl, which would re-save
every part (comments' VML, calc chain, styles tables...).  Instead the uploaded
.xlsx package is copied part-for-part and only the ``<c>`` elements of the
approved target cells are edited: the value changes, the cell's style index
(font, alignment, border, fill, number format) stays exactly as it was.  The
workbook is flagged to recalculate on open so dependent formulas update.
``verify_update`` then proves nothing else changed.
"""
from __future__ import annotations

import datetime as _dt
import io
import posixpath
import re
import zipfile
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple, Union

import openpyxl
from lxml import etree
from openpyxl.utils import column_index_from_string, get_column_letter

from . import values
from .profiles import ColumnRef, SheetLayout

_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_CELL_RE = re.compile(r"^([A-Z]{1,3})(\d+)$")
# Numbers joined by + / - only: a breakup this app wrote earlier, safe to replace.
_BREAKUP_RE = re.compile(r"[\d.]+(?:[+-][\d.]+)*")


def _q(tag: str) -> str:
    return f"{{{_NS}}}{tag}"


class WorkbookError(Exception):
    pass


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------

@dataclass
class CellState:
    value: object
    kind: str   # empty | value | formula | merged


class Workbook:
    """Read-only view of the uploaded workbook."""

    def __init__(self, data: bytes):
        self.data = data
        try:
            self._wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        except Exception as exc:
            raise WorkbookError(f"Could not open the workbook: {exc}") from exc
        self.sheets: List[str] = list(self._wb.sheetnames)
        self._dates: Dict[tuple, Dict[_dt.date, List[int]]] = {}

    def has_sheet(self, sheet: str) -> bool:
        return sheet in self.sheets

    def header(self, sheet: str, row: int, letter: str) -> str:
        value = self._wb[sheet][f"{letter}{row}"].value
        return "" if value is None else str(value)

    def headers(self, sheet: str, row: int) -> List[Tuple[str, str]]:
        ws = self._wb[sheet]
        out = []
        for c in range(1, ws.max_column + 1):
            letter = get_column_letter(c)
            out.append((letter, self.header(sheet, row, letter)))
        return out

    def resolve(self, sheet: str, ref: ColumnRef, header_row: int) -> Tuple[Optional[str], str]:
        """Return ``(letter, problem)``.  A letter is validated against the
        header text saved with it; a header-only reference must be unique."""
        if not ref.is_set:
            return None, "not configured"
        if ref.letter:
            try:
                column_index_from_string(ref.letter)
            except ValueError:
                return None, f"invalid column '{ref.letter}'"
            actual = self.header(sheet, header_row, ref.letter)
            if ref.header and values.label_key(actual) != values.label_key(ref.header):
                return None, (f"column {ref.letter} is now headed '{actual or '(blank)'}', "
                              f"expected '{ref.header}'")
            return ref.letter, ""
        matches = [l for l, text in self.headers(sheet, header_row)
                   if text and values.label_key(text) == values.label_key(ref.header)]
        if len(matches) == 1:
            return matches[0], ""
        if not matches:
            return None, f"no column headed '{ref.header}' in row {header_row}"
        return None, f"header '{ref.header}' repeats in columns {', '.join(matches)}; choose a letter"

    def date_rows(self, sheet: str, letter: str, header_row: int,
                  order: str = values.DMY) -> Dict[_dt.date, List[int]]:
        """Full-column index: normalised date -> worksheet row(s)."""
        key = (sheet, letter, header_row, order)
        if key not in self._dates:
            ws = self._wb[sheet]
            col = column_index_from_string(letter)
            index: Dict[_dt.date, List[int]] = {}
            for r in range(header_row + 1, ws.max_row + 1):
                d = values.parse_date(ws.cell(row=r, column=col).value, order, allow_serial=True)
                if d:
                    index.setdefault(d, []).append(r)
            self._dates[key] = index
        return self._dates[key]

    def cell(self, sheet: str, ref: str) -> CellState:
        ws = self._wb[sheet]
        for rng in ws.merged_cells.ranges:
            if ref in rng and ref != rng.start_cell.coordinate:
                return CellState(None, "merged")
        c = ws[ref]
        if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("=")):
            return CellState(c.value, "formula")
        if c.value is None or (isinstance(c.value, str) and not c.value.strip()):
            return CellState(None, "empty")
        return CellState(c.value, "value")

    def suggest_layout(self, sheet: str) -> SheetLayout:
        """Guess header row and DATE column: the column with the most date
        cells, headed by text in the row just above its first date."""
        ws = self._wb[sheet]
        best = None
        for c in range(1, min(ws.max_column, 30) + 1):
            dates = [r for r in range(1, min(ws.max_row, 400) + 1)
                     if values.parse_date(ws.cell(row=r, column=c).value, allow_serial=False)]
            if len(dates) >= 3 and (best is None or len(dates) > best[1]):
                best = (c, len(dates), dates[0])
        if not best:
            return SheetLayout()
        header_row = max(best[2] - 1, 1)
        letter = get_column_letter(best[0])
        return SheetLayout(header_row, ColumnRef(letter, self.header(sheet, header_row, letter)))


# ---------------------------------------------------------------------------
# Writing: patch target cells inside the original package
# ---------------------------------------------------------------------------

Value = Union[float, int, str]


@dataclass(frozen=True)
class CellChange:
    sheet: str
    cell: str        # e.g. "B272"
    value: Value     # number, or a formula string starting with "="

    @property
    def is_formula(self) -> bool:
        return isinstance(self.value, str) and self.value.startswith("=")


def _sheet_parts(zf: zipfile.ZipFile) -> Dict[str, str]:
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels}
    parts = {}
    for sheet in wb.find(_q("sheets")):
        target = targets[sheet.get(f"{{{_REL_NS}}}id")]
        parts[sheet.get("name")] = (target.lstrip("/") if target.startswith("/")
                                    else posixpath.normpath(posixpath.join("xl", target)))
    return parts


def _parse(xml: bytes):
    return etree.fromstring(xml, etree.XMLParser(remove_blank_text=False, huge_tree=True))


def _serialize(root, original: bytes) -> bytes:
    decl = re.match(rb"<\?xml[^>]*\?>\s*", original)
    head = decl.group(0) if decl else b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
    return head + etree.tostring(root, encoding="UTF-8", xml_declaration=False)


def _number_text(value) -> str:
    f = float(value)
    return str(int(f)) if f.is_integer() and abs(f) < 1e15 else repr(f)


def _col_styles(root) -> Dict[int, str]:
    styles = {}
    cols = root.find(_q("cols"))
    for col in (cols if cols is not None else []):
        if col.get("style") is not None:
            for i in range(int(col.get("min")), int(col.get("max")) + 1):
                styles[i] = col.get("style")
    return styles


def _patch_sheet(xml: bytes, changes: List[CellChange]) -> bytes:
    root = _parse(xml)
    sheet_data = root.find(_q("sheetData"))
    rows = {int(r.get("r")): r for r in sheet_data.findall(_q("row"))}
    col_styles = _col_styles(root)
    for change in changes:
        m = _CELL_RE.match(change.cell)
        letters, row_num = m.group(1), int(m.group(2))
        col_num = column_index_from_string(letters)
        row = rows.get(row_num)
        if row is None:
            row = etree.Element(_q("row"), r=str(row_num))
            later = [r for n, r in sorted(rows.items()) if n > row_num]
            (later[0].addprevious(row) if later else sheet_data.append(row))
            rows[row_num] = row
        cell = next((c for c in row.findall(_q("c")) if c.get("r") == change.cell), None)
        if cell is None:
            cell = etree.Element(_q("c"), r=change.cell)
            style = row.get("s") if row.get("customFormat") in ("1", "true") else col_styles.get(col_num)
            if style is not None:
                cell.set("s", style)
            later = [c for c in row.findall(_q("c"))
                     if column_index_from_string(_CELL_RE.match(c.get("r")).group(1)) > col_num]
            (later[0].addprevious(cell) if later else row.append(cell))
        formula = cell.find(_q("f"))
        if formula is not None and not (formula.get("t") is None and formula.text
                                        and _BREAKUP_RE.fullmatch(formula.text.strip())):
            raise WorkbookError(f"{change.sheet}!{change.cell} holds a formula; refusing to overwrite it.")
        for child in list(cell):
            if child.tag in (_q("v"), _q("is"), _q("f")):
                cell.remove(child)
        cell.attrib.pop("t", None)
        if change.is_formula:
            el = etree.Element(_q("f"))
            el.text = str(change.value)[1:]
        else:
            el = etree.Element(_q("v"))
            el.text = _number_text(change.value)
        cell.insert(0, el)
    return _serialize(root, xml)


_CALCPR_AFTER = ("definedNames", "externalReferences", "functionGroups", "sheets")


def _set_full_calc(xml: bytes) -> bytes:
    root = _parse(xml)
    calc = root.find(_q("calcPr"))
    if calc is None:
        calc = etree.Element(_q("calcPr"))
        anchor = next((root.find(_q(t)) for t in _CALCPR_AFTER if root.find(_q(t)) is not None), None)
        (anchor.addnext(calc) if anchor is not None else root.append(calc))
    calc.set("fullCalcOnLoad", "1")
    return _serialize(root, xml)


def _drop_calc_chain(zf: zipfile.ZipFile, patched: Dict[str, bytes]) -> Optional[str]:
    """Adding new formulas invalidates the calc chain; Excel rebuilds it."""
    if "xl/calcChain.xml" not in zf.namelist():
        return None
    for part, attr in (("[Content_Types].xml", "PartName"), ("xl/_rels/workbook.xml.rels", "Target")):
        xml = patched.get(part, zf.read(part))
        root = _parse(xml)
        for el in list(root):
            if (el.get(attr) or "").lstrip("/").endswith("calcChain.xml"):
                root.remove(el)
        patched[part] = _serialize(root, xml)
    return "xl/calcChain.xml"


def update_workbook(original: bytes, changes: List[CellChange]) -> bytes:
    """Return the uploaded workbook with only ``changes`` applied."""
    if not changes:
        raise WorkbookError("No cell changes to apply.")
    zin = zipfile.ZipFile(io.BytesIO(original))
    parts = _sheet_parts(zin)
    by_part: Dict[str, List[CellChange]] = {}
    for change in changes:
        if change.sheet not in parts:
            raise WorkbookError(f"Worksheet '{change.sheet}' not found.")
        by_part.setdefault(parts[change.sheet], []).append(change)

    patched = {part: _patch_sheet(zin.read(part), chs) for part, chs in by_part.items()}
    patched["xl/workbook.xml"] = _set_full_calc(zin.read("xl/workbook.xml"))
    dropped = _drop_calc_chain(zin, patched) if any(c.is_formula for c in changes) else None

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zout:
        for info in zin.infolist():
            if info.filename == dropped:
                continue
            zout.writestr(info, patched.get(info.filename, zin.read(info.filename)))
    return out.getvalue()


# ---------------------------------------------------------------------------
# Proof that only the approved cells changed
# ---------------------------------------------------------------------------

def _canonical(xml: bytes, drop_cells=frozenset(), drop_calc_flag=False) -> bytes:
    root = _parse(xml)
    sheet_data = root.find(_q("sheetData"))
    if sheet_data is not None:
        for row in list(sheet_data):
            for c in list(row):
                if c.get("r") in drop_cells:
                    row.remove(c)
            # A bare row created only to hold a target cell; applied to both
            # files, so a genuinely pre-existing bare row compares equal too.
            if len(row) == 0 and set(row.attrib) == {"r"}:
                sheet_data.remove(row)
    if drop_calc_flag:
        calc = root.find(_q("calcPr"))
        if calc is not None:
            calc.attrib.pop("fullCalcOnLoad", None)
            if not calc.attrib:
                root.remove(calc)
    return etree.tostring(root, method="c14n")


def _cell_attrs(xml: bytes, refs) -> Dict[str, Optional[str]]:
    root = _parse(xml)
    out = {}
    for c in root.iter(_q("c")):
        if c.get("r") in refs:
            out[c.get("r")] = c.get("s")
    return out


def verify_update(original: bytes, updated: bytes, changes: List[CellChange]) -> List[str]:
    """Empty list = verified: every other package part is byte-identical, the
    edited sheets differ only in the target cells, target cells kept their
    style, and each target holds exactly the planned value."""
    problems: List[str] = []
    zo = zipfile.ZipFile(io.BytesIO(original))
    zu = zipfile.ZipFile(io.BytesIO(updated))
    parts = _sheet_parts(zo)
    formula_mode = any(c.is_formula for c in changes)
    edited = {parts[c.sheet] for c in changes}
    allowed = edited | {"xl/workbook.xml"}
    if formula_mode:
        allowed |= {"[Content_Types].xml", "xl/_rels/workbook.xml.rels"}

    expected_names = set(zo.namelist()) - ({"xl/calcChain.xml"} if formula_mode else set())
    if set(zu.namelist()) != expected_names:
        problems.append("The set of package parts changed.")
    for name in set(zo.namelist()) & set(zu.namelist()):
        if name not in allowed and zo.read(name) != zu.read(name):
            problems.append(f"Unexpected change in {name}.")

    if _canonical(zo.read("xl/workbook.xml"), drop_calc_flag=True) != \
            _canonical(zu.read("xl/workbook.xml"), drop_calc_flag=True):
        problems.append("Workbook settings changed beyond the recalculation flag.")

    for part in edited:
        refs = frozenset(c.cell for c in changes if parts[c.sheet] == part)
        before, after = zo.read(part), zu.read(part)
        if _canonical(before, refs) != _canonical(after, refs):
            problems.append(f"{part}: content other than the target cells changed.")
        styles_before, styles_after = _cell_attrs(before, refs), _cell_attrs(after, refs)
        for ref, style in styles_before.items():
            if styles_after.get(ref) != style:
                problems.append(f"{part}!{ref}: cell style changed.")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(updated), data_only=False)
    except Exception as exc:
        return problems + [f"Updated workbook cannot be opened: {exc}"]
    if wb.sheetnames != openpyxl.load_workbook(io.BytesIO(original)).sheetnames:
        problems.append("Worksheet list changed.")
    for change in changes:
        got = wb[change.sheet][change.cell].value
        if change.is_formula:
            ok = got == change.value
        else:
            ok = isinstance(got, (int, float)) and abs(float(got) - float(change.value)) < 1e-9
        if not ok:
            problems.append(f"{change.sheet}!{change.cell} holds {got!r}, expected {change.value!r}.")
    return problems
